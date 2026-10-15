"""BUG-05 / NV-19: one net-pay calculation everywhere.

Net pay was computed five different ways:
  - the disbursed amount (compute_salary_data_for_month: attendance with
    the company's late/half-day rates, plus incentives, no PF/PT/TDS);
  - the payslip document (CTC structure, loss of pay on a 26-day basis,
    PF/PT/TDS, late and half days paid in full, incentive shown but not
    added);
  - the Excel salary export (its own copy of the first);
  - the portal's "this month" card and /my_payslip_summary (late days
    paid in full, half day fixed at 50%, holidays and leave ignored, a
    flat 12% "PF", overtime added);
  - the mobile app's /api/employee/salary (yet another formula).
An employee's payslip could differ from what was paid.

Owner decision (QA_AUDIT.md §0.1): earnings from attendance under the
company's rules plus incentives -- the amount disbursed until now --
minus the PF/PT/TDS from Payroll settings. Every view now shows that
figure, and it is what the disbursement pays.
"""
import calendar
import datetime
import io
import re

import openpyxl
import pytest

import blueprints.payroll as payroll
from utils.salary_utils import apply_statutory_deductions

PF_CFG = {"pf_employee_pct": 12.0, "pf_employer_pct": 12.0, "professional_tax": 200.0,
          "tds_annual_pct": 0.0, "pf_basic_cap": 15000.0}


def _last_month_working_days():
    first_this = datetime.date.today().replace(day=1)
    last = first_this - datetime.timedelta(days=1)
    days = [datetime.date(last.year, last.month, d) for d in range(1, calendar.monthrange(last.year, last.month)[1] + 1)]
    return last.year, last.month, [d for d in days if d.weekday() != 6]


def _attend(cur, emp, day, att_type):
    cur.execute("INSERT INTO attendance (employee_id, date, login_time, logout_time, status, attendance_type) "
                "VALUES (%s,%s,'09:00','18:00','Login',%s) ON CONFLICT DO NOTHING", (emp, day, att_type))


@pytest.fixture
def pay_month(db_engine, seed_employee):
    """Last month: a paid holiday, an approved leave day, two late days, two
    half days, one absence, full days otherwise, and a 750 incentive."""
    year, month, working = _last_month_working_days()
    emp = seed_employee["employee_id"]
    holiday, leave_day, late, half, absent, full = (working[0], working[1], working[2:4], working[4:6],
                                                    working[6], working[7:])
    cur = db_engine.cursor()
    cur.execute("INSERT INTO salary_config (employee_id, salary_per_day, monthly_ctc, basic_pct) "
                "VALUES (%s, 1000, 26000, 50) ON CONFLICT (employee_id) DO UPDATE "
                "SET salary_per_day=1000, monthly_ctc=26000, basic_pct=50", (emp,))
    cur.execute("UPDATE employees SET email='tst001@test.local' WHERE employee_id=%s", (emp,))
    cur.execute("INSERT INTO holidays (date, name) VALUES (%s, 'QA Holiday') ON CONFLICT DO NOTHING", (holiday,))
    cur.execute("SELECT id FROM leave_types LIMIT 1")
    cur.execute("INSERT INTO leave_requests (employee_id, leave_type_id, leave_date, reason, status) "
                "VALUES (%s,%s,%s,'QA','Approved')", (emp, cur.fetchone()[0], leave_day))
    cur.execute("INSERT INTO incentive_goals (title, incentive_amount) VALUES ('QA goal', 750) RETURNING id")
    cur.execute("INSERT INTO employee_incentives (employee_id, goal_id, month, year, amount) VALUES (%s,%s,%s,%s,750)",
                (emp, cur.fetchone()[0], month, year))
    for d in late:
        _attend(cur, emp, d, "Late - Full Day")
    for d in half:
        _attend(cur, emp, d, "Half Day")
    for d in full:
        _attend(cur, emp, d, "Full Day")
    assert absent not in full
    cur.execute("DELETE FROM payroll_config")
    cur.execute("INSERT INTO payroll_config (pf_employee_pct, pf_employer_pct, professional_tax, tds_annual_pct, "
                "pf_basic_cap) VALUES (12, 12, 200, 0, 15000)")
    cur.close()
    return emp, year, month


def _entry(emp, year, month):
    return payroll.compute_salary_data_for_month(year, month, employee_ids=[emp])[0]


def _slip_net(html):
    m = re.search(r'<div class="amt">Rs\. ([\d,]+\.\d\d)</div>', html)
    assert m, "no net pay on the payslip"
    return float(m.group(1).replace(",", ""))


class TestTheCalculation:
    def test_net_is_attendance_earnings_plus_incentive_minus_pf_pt_tds(self, pay_month):
        emp, year, month = pay_month
        e = _entry(emp, year, month)
        attendance = e["full_earn"] + e["late_earn"] + e["half_earn"]
        assert e["absent"] == 1 and e["late_days"] == 2 and e["half_days"] == 2
        assert e["pf"] == 1560.0          # 12% of min(basic 13,000, cap 15,000)
        assert e["pt"] == 200.0
        assert e["tds"] == 0.0
        assert e["net"] == round(attendance + 750 - 1560 - 200, 2)

    def test_deductions_never_exceed_earnings(self):
        e = apply_statutory_deductions({"spd": 1000, "full_earn": 1000, "late_earn": 0, "half_earn": 0,
                                        "incentive": 0, "monthly_ctc": 26000, "basic_pct": 50,
                                        "gross": 20000}, PF_CFG)
        assert e["net"] == 0.0
        assert e["statutory"] == 1000.0   # PF + PT (1,760) scaled down to the 1,000 earned

    def test_tds_is_monthly_ctc_times_the_annual_rate(self):
        e = apply_statutory_deductions({"spd": 1000, "full_earn": 26000, "monthly_ctc": 26000, "basic_pct": 50,
                                        "gross": 26000}, dict(PF_CFG, tds_annual_pct=5))
        assert e["tds"] == 1300.0


class TestEveryViewShowsThatNet:
    def test_payslip_document(self, client, seed_admin, pay_month):
        emp, year, month = pay_month
        with client.session_transaction() as s:
            s["admin_logged_in"], s["admin_username"], s["admin_role"] = True, seed_admin["username"], "admin"
        html = client.get(f"/view_payslip/{emp}/{year}/{month}").get_data(as_text=True)
        assert _slip_net(html) == _entry(emp, year, month)["net"]

    def test_payslip_email(self, pay_month):
        emp, year, month = pay_month
        err, slip = payroll._payslip_email(emp, year, month)
        assert err is None
        assert _slip_net(slip[2]) == _entry(emp, year, month)["net"]

    def test_excel_export(self, pay_month):
        emp, year, month = pay_month
        buf, _name = payroll._build_salary_report_workbook(year, month)
        ws = openpyxl.load_workbook(io.BytesIO(buf.read())).active
        headers = [c.value for c in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[headers.index("Employee ID")] == emp)
        assert row[headers.index("Net Salary")] == _entry(emp, year, month)["net"]

    def test_my_payslip_summary(self, client, seed_employee, pay_month):
        emp, year, month = pay_month
        with client.session_transaction() as s:
            s["employee_id"] = emp
        data = client.get(f"/my_payslip_summary/{year}/{month}").get_json()
        e = _entry(emp, year, month)
        assert data["net"] == e["net"]
        assert data["statutory"] == e["statutory"]

    def test_mobile_salary_endpoint(self, client, seed_employee, pay_month):
        emp, year, month = pay_month
        token = client.post("/api/employee/login", json={"employee_id": emp,
                                                         "password": seed_employee["password"]}).get_json()["token"]
        data = client.get(f"/api/employee/salary?year={year}&month={month}",
                          headers={"Authorization": f"Bearer {token}"}).get_json()
        assert data["salary"]["net"] == _entry(emp, year, month)["net"]

    def test_portal_card_for_the_current_month(self, client, seed_employee, pay_month, db_engine):
        emp = seed_employee["employee_id"]
        today = datetime.date.today()
        cur = db_engine.cursor()
        for day in range(1, today.day):
            d = today.replace(day=day)
            if d.weekday() != 6:
                _attend(cur, emp, d, "Full Day")
        cur.close()
        with client.session_transaction() as s:
            s["employee_id"] = emp
        html = client.get("/employee_portal").get_data(as_text=True)
        m = re.search(r"Live Estimate</div>\s*<div[^>]*>₹([\d.]+)</div>", html)
        assert m, "no live estimate on the portal"
        assert float(m.group(1)) == _entry(emp, today.year, today.month)["net"]
